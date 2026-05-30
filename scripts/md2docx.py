#!/usr/bin/env python3
"""Markdown → DOCX converter focused on the structures used in our docs.
Headings, tables, fenced code blocks, bullet/numbered lists, bold/italic/inline-code,
blockquotes, and horizontal rules. Tuned for Korean + English mixed documents.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement


CODE_FONT = "Consolas"
BODY_FONT_LATIN = "Calibri"
BODY_FONT_EAST = "Malgun Gothic"
BODY_SIZE = Pt(10.5)
HEADING_SIZES = {1: Pt(20), 2: Pt(16), 3: Pt(13), 4: Pt(11.5)}
GREY_BG = "F2F2F2"


def set_cell_shading(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def style_run(run, *, bold: bool = False, italic: bool = False, code: bool = False,
              size: Pt | None = None) -> None:
    run.bold = bold
    run.italic = italic
    if code:
        run.font.name = CODE_FONT
        # ensure east-asian font also Consolas-like for code consistency
        rPr = run._element.get_or_add_rPr()
        rFonts = OxmlElement("w:rFonts")
        rFonts.set(qn("w:ascii"), CODE_FONT)
        rFonts.set(qn("w:hAnsi"), CODE_FONT)
        rFonts.set(qn("w:eastAsia"), CODE_FONT)
        rPr.append(rFonts)
    else:
        run.font.name = BODY_FONT_LATIN
        rPr = run._element.get_or_add_rPr()
        rFonts = OxmlElement("w:rFonts")
        rFonts.set(qn("w:ascii"), BODY_FONT_LATIN)
        rFonts.set(qn("w:hAnsi"), BODY_FONT_LATIN)
        rFonts.set(qn("w:eastAsia"), BODY_FONT_EAST)
        rPr.append(rFonts)
    if size is not None:
        run.font.size = size
    elif not code:
        run.font.size = BODY_SIZE
    else:
        run.font.size = Pt(9.5)


INLINE_RE = re.compile(
    r"(\*\*[^*]+\*\*)|(\*[^*]+\*)|(`[^`]+`)|(\[[^\]]+\]\([^)]+\))"
)


def add_inline(para, text: str, *, base_bold: bool = False, base_italic: bool = False) -> None:
    """Walk inline markdown spans and add styled runs."""
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            r = para.add_run(text[pos:m.start()])
            style_run(r, bold=base_bold, italic=base_italic)
        token = m.group(0)
        if token.startswith("**") and token.endswith("**"):
            r = para.add_run(token[2:-2])
            style_run(r, bold=True)
        elif token.startswith("*") and token.endswith("*"):
            r = para.add_run(token[1:-1])
            style_run(r, italic=True)
        elif token.startswith("`") and token.endswith("`"):
            r = para.add_run(token[1:-1])
            style_run(r, code=True)
        elif token.startswith("["):
            label, _, _ = token[1:].partition("]")
            r = para.add_run(label)
            style_run(r, italic=True)
        pos = m.end()
    if pos < len(text):
        r = para.add_run(text[pos:])
        style_run(r, bold=base_bold, italic=base_italic)


def add_heading(doc, level: int, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    add_inline(p, text)
    for run in p.runs:
        run.bold = True
        run.font.size = HEADING_SIZES.get(level, Pt(11))
        if level == 1:
            run.font.color.rgb = RGBColor(0x1F, 0x3B, 0x6E)
        elif level == 2:
            run.font.color.rgb = RGBColor(0x2C, 0x55, 0x9E)


def add_paragraph(doc, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    add_inline(p, text)


def add_blockquote(doc, lines: list[str]) -> None:
    text = "\n".join(lines)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(20)
    p.paragraph_format.space_after = Pt(6)
    add_inline(p, text, base_italic=True)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def add_bullet(doc, text: str, *, ordered: bool = False) -> None:
    style = "List Number" if ordered else "List Bullet"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(2)
    add_inline(p, text)


def add_code_block(doc, lines: list[str]) -> None:
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(6)
    para.paragraph_format.left_indent = Pt(10)
    text = "\n".join(lines)
    r = para.add_run(text)
    style_run(r, code=True)
    # subtle grey background via paragraph shading
    pPr = para._element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), GREY_BG)
    pPr.append(shd)


def add_table(doc, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.text = ""
            value = row[j] if j < len(row) else ""
            p = cell.paragraphs[0]
            add_inline(p, value)
            if i == 0:
                set_cell_shading(cell, "DDE6F1")
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph()


def add_hr(doc) -> None:
    p = doc.add_paragraph()
    pPr = p._element.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), "BBBBBB")
    pBdr.append(bottom)
    pPr.append(pBdr)


def parse_table_lines(lines: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for raw in lines:
        s = raw.strip()
        if not s.startswith("|"):
            continue
        # skip separator row (---|---|...)
        if re.fullmatch(r"\|[\s:-]+(\|[\s:-]+)+\|?", s):
            continue
        cells = [c.strip() for c in s.strip("|").split("|")]
        rows.append(cells)
    return rows


def render(md_text: str, out_path: Path) -> None:
    doc = Document()
    # default style: set both Latin and East-Asian fonts so Korean renders
    style = doc.styles["Normal"]
    style.font.name = BODY_FONT_LATIN
    style.font.size = BODY_SIZE
    rpr = style.element.get_or_add_rPr()
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), BODY_FONT_LATIN)
    rFonts.set(qn("w:hAnsi"), BODY_FONT_LATIN)
    rFonts.set(qn("w:eastAsia"), BODY_FONT_EAST)
    rpr.append(rFonts)

    lines = md_text.splitlines()
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.rstrip()

        # fenced code block
        if stripped.startswith("```"):
            j = i + 1
            buf: list[str] = []
            while j < n and not lines[j].startswith("```"):
                buf.append(lines[j])
                j += 1
            add_code_block(doc, buf)
            i = j + 1
            continue

        # table block
        if stripped.startswith("|"):
            j = i
            block: list[str] = []
            while j < n and lines[j].strip().startswith("|"):
                block.append(lines[j])
                j += 1
            rows = parse_table_lines(block)
            add_table(doc, rows)
            i = j
            continue

        # blockquote block
        if stripped.startswith(">"):
            j = i
            block_lines: list[str] = []
            while j < n and lines[j].strip().startswith(">"):
                block_lines.append(lines[j].strip()[1:].strip())
                j += 1
            add_blockquote(doc, block_lines)
            i = j
            continue

        # horizontal rule
        if re.fullmatch(r"-{3,}|_{3,}|\*{3,}", stripped):
            add_hr(doc)
            i += 1
            continue

        # heading
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            add_heading(doc, level, m.group(2).strip())
            i += 1
            continue

        # numbered list
        m = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if m:
            add_bullet(doc, m.group(1).strip(), ordered=True)
            i += 1
            continue

        # bullet list
        m = re.match(r"^\s*[-*]\s+(.*)$", line)
        if m:
            add_bullet(doc, m.group(1).strip())
            i += 1
            continue

        # blank line
        if not stripped:
            i += 1
            continue

        # default paragraph
        add_paragraph(doc, stripped)
        i += 1

    doc.save(out_path)


def main() -> None:
    if len(sys.argv) != 3:
        print("usage: md2docx.py <input.md> <output.docx>", file=sys.stderr)
        sys.exit(2)
    src = Path(sys.argv[1]).read_text(encoding="utf-8")
    out = Path(sys.argv[2])
    render(src, out)
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
